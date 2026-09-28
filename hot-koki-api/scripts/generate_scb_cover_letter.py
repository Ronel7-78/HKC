from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "Lettre_motivation_SCB_Cameroun.docx"


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


document = Document()
section = document.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.4)

styles = document.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.08

table = document.add_table(rows=1, cols=2)
table.autofit = False
table.columns[0].width = Cm(8)
table.columns[1].width = Cm(8)
left, right = table.rows[0].cells
set_cell_margins(left)
set_cell_margins(right)

p = left.paragraphs[0]
p.paragraph_format.space_after = Pt(0)
r = p.add_run("[NOM ET PRÉNOM]")
r.bold = True
r.font.size = Pt(12)
r.font.color.rgb = RGBColor(31, 53, 36)
for line in ("[Téléphone]", "[Adresse e-mail]", "[Ville]"):
    p.add_run("\n" + line)

p = right.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.paragraph_format.space_after = Pt(0)
r = p.add_run("À l’attention de la Direction\ndes Ressources Humaines\n")
r.bold = True
p.add_run("SCB Cameroun\nYaoundé")

document.add_paragraph()

p = document.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.add_run("Yaoundé, le [date]")

p = document.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(14)
r = p.add_run("Objet : Candidature spontanée – Systèmes d’information, réseaux et digital")
r.bold = True
r.font.color.rgb = RGBColor(31, 53, 36)

paragraphs = [
    "Madame, Monsieur,",
    (
        "Administrateur systèmes et réseaux, avec une expérience pratique dans le déploiement "
        "d’infrastructures, la sécurisation des environnements informatiques et le développement "
        "de solutions numériques, je souhaite vous soumettre ma candidature spontanée afin "
        "d’intégrer les équipes de SCB Cameroun."
    ),
    (
        "Mon parcours m’a permis de développer des compétences en administration de serveurs "
        "Linux, configuration réseau, sécurisation des accès, gestion de bases de données et "
        "déploiement d’applications. J’ai notamment participé à la conception et à la mise en "
        "production d’une plateforme mobile complète intégrant une API sécurisée, des services "
        "de paiement, des notifications et des environnements distincts de développement, de test "
        "et de production."
    ),
    (
        "Rigoureux, autonome et attentif aux questions de disponibilité, de confidentialité et de "
        "traçabilité, je souhaite mettre mes compétences au service d’une institution bancaire "
        "engagée dans la modernisation de ses infrastructures et la transformation digitale. "
        "Rejoindre SCB Cameroun représenterait pour moi l’occasion de contribuer à des projets "
        "exigeants tout en consolidant mon expérience dans un environnement où la sécurité et la "
        "continuité des services sont essentielles."
    ),
    (
        "Disponible et motivé, je serais heureux de vous rencontrer afin de vous présenter plus "
        "précisément mon parcours et les contributions que je pourrais apporter à vos équipes."
    ),
    (
        "Je vous prie d’agréer, Madame, Monsieur, l’expression de ma considération distinguée."
    ),
]

for index, text in enumerate(paragraphs):
    p = document.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if index == 0:
        p.paragraph_format.space_after = Pt(10)

document.add_paragraph()
p = document.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("[NOM ET PRÉNOM]")
r.bold = True

document.save(OUTPUT)
print(OUTPUT)

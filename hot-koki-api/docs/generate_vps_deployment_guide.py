from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Guide_deploiement_Hot_Koki_VPS.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.1)
section.right_margin = Cm(2.1)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(34)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(29, 67, 43)
for name, size, color in [
    ("Heading 1", 20, RGBColor(29, 67, 43)),
    ("Heading 2", 15, RGBColor(224, 70, 18)),
    ("Heading 3", 12, RGBColor(29, 67, 43)),
]:
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = color
    styles[name].paragraph_format.space_before = Pt(10)
    styles[name].paragraph_format.space_after = Pt(5)

code_style = styles.add_style("HKC Code", WD_STYLE_TYPE.PARAGRAPH)
code_style.font.name = "Consolas"
code_style.font.size = Pt(8.5)
code_style.font.color.rgb = RGBColor(35, 35, 35)
code_style.paragraph_format.left_indent = Cm(0.35)
code_style.paragraph_format.right_indent = Cm(0.35)
code_style.paragraph_format.space_before = Pt(3)
code_style.paragraph_format.space_after = Pt(7)

note_style = styles.add_style("HKC Note", WD_STYLE_TYPE.PARAGRAPH)
note_style.font.name = "Aptos"
note_style.font.size = Pt(9.5)
note_style.font.italic = True
note_style.font.color.rgb = RGBColor(84, 70, 60)
note_style.paragraph_format.left_indent = Cm(0.4)
note_style.paragraph_format.space_after = Pt(7)


def code(text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.2)
    shade(cell, "F3F1ED")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.style = "HKC Code"
    p.add_run(text.strip())
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def note(text):
    p = doc.add_paragraph(style="HKC Note")
    p.add_run("Important — ").bold = True
    p.add_run(text)


def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


# Couverture
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("HOT KOKI CHAUD")
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(224, 70, 18)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Guide complet de déploiement\nde l’API sur VPS")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Laravel 13 • Ubuntu 24.04 • Nginx • MySQL • HTTPS • Flutter")
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(70, 70, 70)

doc.add_paragraph()
table = doc.add_table(rows=5, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Light Shading Accent 1"
facts = [
    ("Environnement", "Staging / préproduction"),
    ("API", "https://api-staging.hotkokichaud.digital"),
    ("Dépôt", "https://github.com/Ronel7-78/HKC.git"),
    ("Backend", "hot-koki-api"),
    ("Mise à jour", "14 septembre 2026"),
]
for row, (key, value) in zip(table.rows, facts):
    row.cells[0].text = key
    row.cells[1].text = value

doc.add_paragraph()
note("Ce document ne contient aucun secret. Les valeurs telles que IP_DU_VPS, TON_EMAIL et TON_MOT_DE_PASSE doivent être remplacées localement. Ne jamais envoyer une clé SSH privée, un mot de passe ou une clé d’API dans une messagerie.")

doc.add_page_break()
doc.add_heading("Sommaire", level=1)
p = doc.add_paragraph()
add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
doc.add_paragraph("Dans Microsoft Word : clic droit sur le sommaire → Mettre à jour les champs.", style="HKC Note")

doc.add_page_break()
doc.add_heading("1. Architecture retenue", level=1)
doc.add_paragraph("Le premier déploiement est un environnement de staging stable. Il permet de tester l’application mobile et les intégrations sans confondre les essais avec la future production.")
table = doc.add_table(rows=1, cols=3)
table.style = "Light Shading Accent 1"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Usage", "Adresse", "Rôle"]
for i, value in enumerate(headers):
    table.rows[0].cells[i].text = value
    shade(table.rows[0].cells[i], "1D432B")
    for run in table.rows[0].cells[i].paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
set_repeat_table_header(table.rows[0])
for values in [
    ("Staging", "api-staging.hotkokichaud.digital", "Tests actuels"),
    ("Production", "api.hotkokichaud.digital", "À réserver pour la mise en ligne"),
]:
    cells = table.add_row().cells
    for i, value in enumerate(values):
        cells[i].text = value

bullets([
    "Nginx expose uniquement le dossier public de Laravel.",
    "PHP-FPM 8.4 exécute Laravel 13 et les dépendances Symfony 8.1.",
    "MySQL utilise une base et un utilisateur applicatif dédiés.",
    "Supervisor maintient les workers de paiement actifs.",
    "Cron lance le scheduler Laravel chaque minute pour le polling des paiements.",
    "Certbot fournit et renouvelle automatiquement le certificat HTTPS.",
])

doc.add_heading("2. Prérequis", level=1)
bullets([
    "Un VPS Ubuntu 24.04 LTS avec une IPv4 publique.",
    "Un accès SSH root initial.",
    "Le domaine hotkokichaud.digital administré chez Hostinger.",
    "Le dépôt GitHub public de l’application.",
    "Les secrets MySQL et MTN conservés dans un gestionnaire de mots de passe.",
])

doc.add_heading("3. Configurer le DNS Hostinger", level=1)
doc.add_paragraph("Dans hPanel : Domaines → hotkokichaud.digital → DNS / Serveurs de noms. Ajouter l’enregistrement suivant sans modifier les enregistrements @ et www existants.")
table = doc.add_table(rows=2, cols=4)
table.style = "Light Shading Accent 1"
for i, value in enumerate(["Type", "Nom", "Valeur", "TTL"]):
    table.rows[0].cells[i].text = value
for i, value in enumerate(["A", "api-staging", "IPv4 publique du VPS", "300 ou Auto"]):
    table.rows[1].cells[i].text = value
code("dig +short api-staging.hotkokichaud.digital")
doc.add_paragraph("Le résultat doit être exactement l’IPv4 publique du VPS.")

doc.add_heading("4. Connexion et préparation d’Ubuntu", level=1)
code("ssh root@IP_DU_VPS")
doc.add_paragraph("Si le port n’est pas 22 :")
code("ssh -p NUMERO_DU_PORT root@IP_DU_VPS")
code("apt update\napt upgrade -y\ntimedatectl set-timezone Africa/Douala\ncat /etc/os-release")

doc.add_heading("5. Installer la pile serveur", level=1)
code("apt install -y nginx mysql-server git curl unzip supervisor certbot python3-certbot-nginx software-properties-common")
doc.add_paragraph("Le projet verrouille Symfony 8.1, qui exige PHP 8.4.1 ou supérieur. PHP 8.3 d’Ubuntu ne suffit donc pas pour ce lockfile.")
code("add-apt-repository ppa:ondrej/php -y\napt update")
code("apt install -y php8.4-fpm php8.4-cli php8.4-mysql php8.4-curl php8.4-mbstring php8.4-xml php8.4-bcmath php8.4-zip php8.4-intl php8.4-gd")
code("update-alternatives --set php /usr/bin/php8.4\nsystemctl enable --now nginx mysql php8.4-fpm supervisor\nphp -v")
note("Ne pas exécuter composer update sur le VPS pour contourner une incompatibilité PHP. Cette commande modifierait les dépendances verrouillées du projet.")

doc.add_heading("6. Installer Composer", level=1)
code("cd /tmp\ncurl -sS https://getcomposer.org/installer -o composer-setup.php\nphp composer-setup.php --install-dir=/usr/local/bin --filename=composer\ncomposer --version")

doc.add_heading("7. Pare-feu et utilisateur de déploiement", level=1)
code("ufw allow OpenSSH\nufw allow 'Nginx Full'\nufw enable\nufw status")
note("Toujours autoriser OpenSSH avant d’activer UFW afin de ne pas perdre l’accès distant.")
code("adduser --disabled-password --gecos \"\" deploy\nusermod -aG www-data deploy\nid deploy")

doc.add_heading("8. Créer la base MySQL", level=1)
doc.add_paragraph("Générer un mot de passe fort, puis le conserver hors du serveur dans un gestionnaire de mots de passe.")
code("openssl rand -base64 32\nmysql")
code("CREATE DATABASE hot_koki_staging\nCHARACTER SET utf8mb4\nCOLLATE utf8mb4_unicode_ci;\n\nCREATE USER 'hot_koki_app'@'localhost'\nIDENTIFIED BY 'TON_MOT_DE_PASSE';\n\nGRANT SELECT, INSERT, UPDATE, DELETE, CREATE, ALTER, INDEX, DROP, REFERENCES\nON hot_koki_staging.*\nTO 'hot_koki_app'@'localhost';\n\nFLUSH PRIVILEGES;\nEXIT;")
code("mysql -u hot_koki_app -p hot_koki_staging")

doc.add_heading("9. Cloner et installer le backend", level=1)
code("mkdir -p /var/www/hot-koki\ngit clone https://github.com/Ronel7-78/HKC.git /var/www/hot-koki/source\nchown -R deploy:www-data /var/www/hot-koki")
code("cd /var/www/hot-koki/source/hot-koki-api\nrunuser -u deploy -- composer install --no-dev --optimize-autoloader\nrunuser -u deploy -- php artisan --version")
note("Composer ne doit pas être exécuté en root. Si « deploy: no such user » apparaît, reprendre l’étape 7.")

doc.add_heading("10. Créer et sécuriser le .env", level=1)
code("cd /var/www/hot-koki/source/hot-koki-api\nrunuser -u deploy -- cp .env.example .env\nnano .env")
doc.add_paragraph("Valeurs minimales de staging :")
code('''APP_NAME="Hot Koki Chaud"
APP_ENV=staging
APP_KEY=
APP_DEBUG=false
APP_URL=https://api-staging.hotkokichaud.digital
APP_TIMEZONE=Africa/Douala
APP_LOCALE=fr
APP_FALLBACK_LOCALE=fr
EMAIL_VERIFICATION_ENABLED=false
LOG_LEVEL=warning

DB_CONNECTION=mysql
DB_HOST=127.0.0.1
DB_PORT=3306
DB_DATABASE=hot_koki_staging
DB_USERNAME=hot_koki_app
DB_PASSWORD="TON_MOT_DE_PASSE_MYSQL"

SESSION_DRIVER=database
CACHE_STORE=database
QUEUE_CONNECTION=database
FILESYSTEM_DISK=public

MAIL_MAILER=log
MAIL_FROM_ADDRESS="no-reply@hotkokichaud.digital"
MAIL_FROM_NAME="Hot Koki Chaud"

ORANGE_MONEY_ENABLED=false''')
doc.add_paragraph("Configuration MTN sandbox :")
code('''MTN_MOMO_BASE_URL=https://sandbox.momodeveloper.mtn.com
MTN_MOMO_TARGET_ENVIRONMENT=sandbox
MTN_MOMO_CURRENCY=EUR
MTN_MOMO_CALLBACK_BASE_URL=https://api-staging.hotkokichaud.digital
MTN_MOMO_SUBSCRIPTION_KEY="TA_PRIMARY_KEY"
MTN_MOMO_API_USER="TON_API_USER"
MTN_MOMO_API_KEY="TON_API_KEY"
MTN_MOMO_CALLBACK_ALLOWED_IPS=
MTN_MOMO_POLL_MAX_ATTEMPTS=8''')
note("Le vrai .env reste uniquement sur le VPS, est ignoré par Git et ne doit jamais être copié dans un ticket, un chat ou un dépôt.")
code("chown deploy:www-data .env\nchmod 640 .env\nrunuser -u deploy -- php artisan key:generate")

doc.add_heading("11. Initialiser Laravel", level=1)
code("runuser -u deploy -- php artisan migrate --force\nrunuser -u deploy -- php artisan storage:link\nchown -R deploy:www-data storage bootstrap/cache\nchmod -R 775 storage bootstrap/cache\nrunuser -u deploy -- php artisan about\nrunuser -u deploy -- php artisan mtn:check-env")

doc.add_heading("12. Configurer Nginx", level=1)
code("nano /etc/nginx/sites-available/hot-koki-staging")
code('''server {
    listen 80;
    listen [::]:80;
    server_name api-staging.hotkokichaud.digital;
    root /var/www/hot-koki/source/hot-koki-api/public;

    index index.php;
    charset utf-8;
    client_max_body_size 12M;

    add_header X-Content-Type-Options "nosniff" always;
    add_header X-Frame-Options "SAMEORIGIN" always;
    add_header Referrer-Policy "strict-origin-when-cross-origin" always;

    location / {
        try_files $uri $uri/ /index.php?$query_string;
    }

    location = /favicon.ico { access_log off; log_not_found off; }
    location = /robots.txt { access_log off; log_not_found off; }

    location ~ ^/index\\.php(/|$) {
        fastcgi_pass unix:/run/php/php8.4-fpm.sock;
        fastcgi_param SCRIPT_FILENAME $realpath_root$fastcgi_script_name;
        include fastcgi_params;
        fastcgi_hide_header X-Powered-By;
        fastcgi_read_timeout 60;
    }

    location ~ \\.php$ { return 404; }
    location ~ /\\.(?!well-known).* { deny all; }
}''')
code("ln -s /etc/nginx/sites-available/hot-koki-staging /etc/nginx/sites-enabled/hot-koki-staging\nrm /etc/nginx/sites-enabled/default\nnginx -t\nsystemctl reload nginx\ncurl -I http://api-staging.hotkokichaud.digital")

doc.add_heading("13. Activer HTTPS", level=1)
code("certbot --nginx -d api-staging.hotkokichaud.digital --email TON_EMAIL --agree-tos --no-eff-email --redirect")
code("curl -I https://api-staging.hotkokichaud.digital\ncertbot renew --dry-run")

doc.add_heading("14. Workers et polling automatique", level=1)
code("nano /etc/supervisor/conf.d/hot-koki-worker.conf")
code('''[program:hot-koki-worker]
process_name=%(program_name)s_%(process_num)02d
command=/usr/bin/php8.4 /var/www/hot-koki/source/hot-koki-api/artisan queue:work database --queue=paiements,default --sleep=3 --tries=3 --timeout=60 --max-time=3600
directory=/var/www/hot-koki/source/hot-koki-api
autostart=true
autorestart=true
stopasgroup=true
killasgroup=true
user=deploy
numprocs=2
redirect_stderr=true
stdout_logfile=/var/www/hot-koki/source/hot-koki-api/storage/logs/worker.log
stopwaitsecs=90''')
code("supervisorctl reread\nsupervisorctl update\nsupervisorctl start hot-koki-worker:*\nsupervisorctl status")
code("nano /etc/cron.d/hot-koki-scheduler")
code("* * * * * deploy cd /var/www/hot-koki/source/hot-koki-api && /usr/bin/php8.4 artisan schedule:run > /dev/null 2>&1")
code("chmod 644 /etc/cron.d/hot-koki-scheduler\nsystemctl enable --now cron\nrunuser -u deploy -- php artisan schedule:list")

doc.add_heading("15. Créer le premier administrateur", level=1)
doc.add_paragraph("Le seeder AdminSeeder est réservé au développement local. Sur staging, utiliser Tinker avec un prompt de mot de passe masqué.")
code("cd /var/www/hot-koki/source/hot-koki-api\nrunuser -u deploy -- php artisan tinker")
code('''$email = Laravel\\Prompts\\text(label: 'Email administrateur', required: true);
$telephone = Laravel\\Prompts\\text(label: 'Téléphone administrateur', required: true);
$password = Laravel\\Prompts\\password(label: 'Mot de passe administrateur', required: true);
$user = App\\Models\\User::create(['name' => 'Administrateur Hot Koki', 'email' => mb_strtolower(trim($email)), 'telephone' => $telephone, 'password' => $password]);
$user->forceFill(['role' => 'admin', 'email_verified_at' => now()])->save();
App\\Models\\Admin::create(['user_id' => $user->id, 'nom' => 'Administrateur', 'prenom' => 'Hot Koki']);
$user->fresh()->load('admin');
unset($password);
exit''')

doc.add_heading("16. Connecter Flutter au staging", level=1)
code('''cd "/home/mrroot/Bureau/HKC APP/hot_koki_app"
flutter clean
flutter pub get
flutter devices
flutter run --flavor staging -d R5CX20P9LLT --dart-define=API_BASE_URL=https://api-staging.hotkokichaud.digital/api''')
note("Ne pas copier une URL transformée en lien Markdown. La variable s’écrit exactement API_BASE_URL, sans antislash. Le flavor staging est obligatoire car le projet génère des APK development, staging et production.")
doc.add_paragraph("APK staging produit :")
code("build/app/outputs/flutter-apk/app-staging-debug.apk")

doc.add_heading("17. Mettre l’API à jour", level=1)
doc.add_paragraph("Avant chaque mise à jour importante, sauvegarder la base. Puis :")
code('''cd /var/www/hot-koki/source
runuser -u deploy -- git pull --ff-only origin main
cd hot-koki-api
runuser -u deploy -- composer install --no-dev --optimize-autoloader
runuser -u deploy -- php artisan down --retry=60
runuser -u deploy -- php artisan migrate --force
runuser -u deploy -- php artisan optimize:clear
runuser -u deploy -- php artisan optimize
runuser -u deploy -- php artisan up
supervisorctl restart hot-koki-worker:*''')
doc.add_paragraph("Une modification purement interne au backend ne demande pas de reconstruire Flutter. Une nouvelle compilation mobile est nécessaire si l’URL ou le contrat de l’API change.")

doc.add_heading("18. Sauvegardes", level=1)
doc.add_heading("Sauvegarder MySQL", level=2)
code("mkdir -p /var/backups/hot-koki\nmysqldump --single-transaction -u hot_koki_app -p hot_koki_staging | gzip > /var/backups/hot-koki/hot_koki_staging_DATE.sql.gz")
doc.add_heading("Sauvegarder les fichiers téléversés", level=2)
code("tar -czf /var/backups/hot-koki/storage_DATE.tar.gz -C /var/www/hot-koki/source/hot-koki-api storage/app/public")
note("Automatiser ensuite ces sauvegardes, les chiffrer et en conserver une copie hors du VPS. Tester périodiquement la restauration.")

doc.add_heading("19. MTN MoMo sur le domaine stable", level=1)
doc.add_paragraph("Le providerCallbackHost de l’utilisateur API sandbox doit correspondre exactement à :")
code("api-staging.hotkokichaud.digital")
doc.add_paragraph("Si l’utilisateur sandbox a été créé avec une ancienne adresse trycloudflare.com, le recréer avec le domaine stable, générer une nouvelle api_key puis remplacer les valeurs uniquement dans le .env du VPS.")
code("cd /var/www/hot-koki/source/hot-koki-api\nrunuser -u deploy -- php artisan optimize:clear\nrunuser -u deploy -- php artisan mtn:check-env\nrunuser -u deploy -- php artisan mtn:test-config\nrunuser -u deploy -- php artisan optimize\nsupervisorctl restart hot-koki-worker:*")
doc.add_paragraph("Le callback généré par l’application suit cette forme :")
code("https://api-staging.hotkokichaud.digital/api/webhooks/mtn-momo/{transactionHash}")

doc.add_heading("20. Vérifications finales", level=1)
code("systemctl is-active nginx php8.4-fpm mysql supervisor cron\nsupervisorctl status\nnginx -t\ncurl -I https://api-staging.hotkokichaud.digital\ncurl -i https://api-staging.hotkokichaud.digital/api/catalogue")
bullets([
    "APP_DEBUG=false.",
    "Le .env appartient à deploy:www-data et possède les permissions 640.",
    "Le compte MySQL applicatif n’est pas root.",
    "Les workers Supervisor sont RUNNING.",
    "Le certificat HTTPS et son renouvellement sont valides.",
    "Les sauvegardes de base et de fichiers sont opérationnelles.",
    "Les clés sandbox et production restent distinctes.",
])

doc.add_heading("21. Dépannage rapide", level=1)
table = doc.add_table(rows=1, cols=3)
table.style = "Light Shading Accent 1"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, value in enumerate(["Symptôme", "Cause probable", "Correction"]):
    table.rows[0].cells[i].text = value
    shade(table.rows[0].cells[i], "1D432B")
    for run in table.rows[0].cells[i].paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
set_repeat_table_header(table.rows[0])
rows = [
    ("Composer exige PHP >= 8.4.1", "Symfony 8.1 verrouillé", "Installer PHP 8.4, ne pas faire composer update"),
    ("deploy: no such user", "Utilisateur absent", "Créer deploy puis l’ajouter au groupe www-data"),
    ("502 Bad Gateway", "Socket PHP incorrect ou FPM arrêté", "Vérifier php8.4-fpm et /run/php/php8.4-fpm.sock"),
    ("vendor/autoload.php absent", "composer install a échoué", "Corriger PHP puis relancer Composer avec deploy"),
    ("Flutter ne trouve pas l’APK", "Flavor non précisé", "Ajouter --flavor staging"),
    ("Catalogue hors ligne", "URL Flutter erronée", "Utiliser HTTPS et terminer l’URL par /api"),
    ("Callback MTN refusé", "providerCallbackHost différent", "Recréer l’utilisateur sandbox avec le domaine stable"),
    ("Paiement reste en attente", "Worker ou scheduler arrêté", "Vérifier Supervisor, cron et les logs Laravel"),
]
for values in rows:
    cells = table.add_row().cells
    for i, value in enumerate(values):
        cells[i].text = value

doc.add_heading("Commandes de diagnostic", level=2)
code('''journalctl -u nginx --since "30 minutes ago"
journalctl -u php8.4-fpm --since "30 minutes ago"
tail -n 100 /var/www/hot-koki/source/hot-koki-api/storage/logs/laravel.log
tail -n 100 /var/www/hot-koki/source/hot-koki-api/storage/logs/worker.log
supervisorctl status
runuser -u deploy -- php artisan queue:failed''')
note("Les logs ne doivent jamais contenir de token d’accès, de mot de passe, de clé API ou de MSISDN complet.")

doc.add_heading("22. Passage futur en production", level=1)
numbered([
    "Créer api.hotkokichaud.digital et un environnement serveur distinct si possible.",
    "Créer une base de production distincte.",
    "Utiliser APP_ENV=production et APP_DEBUG=false.",
    "Configurer les identifiants MTN et Orange de production, jamais ceux du sandbox.",
    "Configurer un vrai SMTP et réactiver la vérification email.",
    "Mettre en place les listes blanches IP exigées par les opérateurs.",
    "Exécuter les tests d’intégration, l’audit de sécurité et les tests de restauration.",
    "Construire Flutter avec le flavor production et l’URL de production.",
])
code("flutter build appbundle --flavor production --dart-define=API_BASE_URL=https://api.hotkokichaud.digital/api")

# Pied de page
for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Hot Koki Chaud — Guide de déploiement VPS — ")
    add_field(footer, "PAGE")

doc.core_properties.title = "Guide complet de déploiement Hot Koki sur VPS"
doc.core_properties.subject = "Déploiement Laravel 13 sur Ubuntu 24.04 et connexion Flutter"
doc.core_properties.author = "Hot Koki Chaud"
doc.core_properties.keywords = "Laravel, VPS, Ubuntu, Nginx, MySQL, Flutter, MTN MoMo"
doc.save(OUTPUT)
print(OUTPUT)
